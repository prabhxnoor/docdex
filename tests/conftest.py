from __future__ import annotations

import sys
from pathlib import Path

import pytest

SRC = Path(__file__).resolve().parents[1] / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from docdex.config import Project, ensure_state_dirs  # noqa: E402


def make_pdf(text: str) -> bytes:
    """Build a minimal but structurally valid one-page PDF."""
    stream = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
         b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"),
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF\n").encode()
    return bytes(out)


def make_docx(path: Path, text: str) -> None:
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


def make_xlsx(path: Path, cells) -> None:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in cells:
        ws.append(row)
    wb.save(str(path))


def make_encrypted_pdf(text: str, password: str) -> bytes:
    """An RC4-128 password-encrypted one-page PDF wrapping make_pdf(text)."""
    import io
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(make_pdf(text))).pages:
        writer.add_page(page)
    writer.encrypt(user_password=password, algorithm="RC4-128")
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


@pytest.fixture
def corpus(tmp_path: Path) -> Path:
    """A small synthetic document corpus (no docdex project yet)."""
    root = tmp_path / "proj"
    (root / "Reports").mkdir(parents=True)
    (root / "Notes" / "Deep" / "Nested").mkdir(parents=True)
    (root / "node_modules").mkdir()
    (root / "_indexes").mkdir()  # sibling that must NOT be confused with _index

    (root / "Reports" / "Q1 report.md").write_text(
        "# Q1 report\nquarterly revenue grew ZEPHYRTOKEN strongly\n", encoding="utf-8")
    (root / "Reports" / "data.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    (root / "Notes" / "hello.txt").write_text("plain hello world note\n", encoding="utf-8")
    (root / "Notes" / "Deep" / "Nested" / "deep file.md").write_text(
        "nested unicode naïve café content\n", encoding="utf-8")
    (root / "_indexes" / "sibling.md").write_text("sibling folder content\n", encoding="utf-8")
    (root / "unsupported.bin").write_bytes(b"\x00\x01\x02binary")
    (root / "node_modules" / "junk.txt").write_text("never index me\n", encoding="utf-8")
    (root / ".hiddenfile.md").write_text("hidden\n", encoding="utf-8")
    (root / "~$lock.docx").write_text("office lock\n", encoding="utf-8")

    make_docx(root / "Reports" / "board deck.docx", "docx body XANTHICWORD inside")
    make_xlsx(root / "Reports" / "numbers.xlsx", [["h1", "h2"], ["v1", 42]])
    (root / "Reports" / "scan.pdf").write_bytes(make_pdf(""))  # no extractable text
    pdf_text = ("QUARKPDF quantum key distribution content lives here. " * 8).strip()
    (root / "Notes" / "real.pdf").write_bytes(make_pdf(pdf_text))
    (root / "diagram.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 64)
    return root


@pytest.fixture
def project(corpus: Path) -> Project:
    """An initialized (but not yet synced) docdex project over the corpus."""
    p = Project.create(corpus, index_dir="_index")
    ensure_state_dirs(p)
    p.save()
    return p


@pytest.fixture
def synced(project: Project) -> Project:
    from docdex.sync import run_sync
    run_sync(project, quiet=True)
    return project
